"""Build a readable Word export of the current SNO Level-1/Level-2 directory."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data/processed/record_role_classification_v1/targeted_name_repairs/name_only_targeted_repair_v1"
SOURCE = DATA_DIR / "sno_level_1_2_directory_v1.json"
OUTPUT = DATA_DIR / "SNO一级二级目录结构_v1.docx"

NAVY = "18344A"
PALE_BLUE = "EDF4F8"
GRID = "D9D9D9"
GRAY = "666666"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "000000", size: int = 10, align=None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_border(cell)


def style_table(table) -> None:
    table.autofit = False
    table.style = "Table Grid"
    table.rows[0].height = None
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)


def add_table(document: Document, entries: list[dict]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.columns[0].width = Inches(0.45)
    table.columns[1].width = Inches(2.15)
    table.columns[2].width = Inches(0.85)
    table.columns[3].width = Inches(2.75)
    headers = ("序号", "二级目录", "三级分类数", "稳定 ID")
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True, color="FFFFFF", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    style_table(table)
    for index, entry in enumerate(entries, 1):
        row = table.add_row()
        if index % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
        name = entry["display_name"]
        if entry["canonical_name"] != name:
            name = f"{name}\n（正式名称：{entry['canonical_name']}）"
        set_cell_text(row.cells[0], index, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], name, bold=True, size=10)
        set_cell_text(row.cells[2], entry["level_3_category_count"], size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[3], entry["level_2_id"], color=GRAY, size=8)
    document.add_paragraph().paragraph_format.space_after = Pt(3)


def set_document_defaults(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("SNO 职业目录 一级和二级结构")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite {OUTPUT}")
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    document = Document()
    set_document_defaults(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("SNO 一级二级职业目录")
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.add_run("用于审阅当前 SNO 展示层的一级和二级目录结构").italic = True
    summary = data["summary"]
    intro = document.add_paragraph()
    intro.add_run("目录范围：").bold = True
    intro.add_run(f"正式目录共 {summary['formal_level_1_count']} 个一级、{summary['formal_level_2_count']} 个二级。")
    intro.add_run("下文的展示层暂存分组不属于冻结的正式目录，用于承接尚不足以进入正式细分类别的岗位。")

    document.add_heading("正式一级和二级目录", level=1)
    for root in data["formal_taxonomy"]:
        heading = document.add_heading(root["level_1_name"], level=2)
        heading.paragraph_format.keep_with_next = True
        add_table(document, root["level_2"])

    document.add_page_break()
    document.add_heading("展示层暂存分组", level=1)
    note = document.add_paragraph("以下分组仅用于前端展示和兜底，不改变正式一级、二级职业目录。")
    note.paragraph_format.space_after = Pt(8)
    for root in data["display_only_holding_groups"]:
        heading = document.add_heading(root["level_1_name"], level=2)
        heading.paragraph_format.keep_with_next = True
        if root["level_2"]:
            table = document.add_table(rows=1, cols=2)
            table.columns[0].width = Inches(0.55)
            table.columns[1].width = Inches(5.65)
            for cell, text in zip(table.rows[0].cells, ("序号", "展示层二级分组")):
                set_cell_text(cell, text, bold=True, color="FFFFFF", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            style_table(table)
            for index, child in enumerate(root["level_2"], 1):
                row = table.add_row()
                if index % 2 == 0:
                    for cell in row.cells:
                        set_cell_shading(cell, PALE_BLUE)
                set_cell_text(row.cells[0], index, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                set_cell_text(row.cells[1], child["display_name"], size=10)
        else:
            paragraph = document.add_paragraph("该分组直接连接真实岗位记录，不再设置二级展示节点。")
            paragraph.paragraph_format.left_indent = Inches(0.15)

    temp = OUTPUT.with_suffix(".tmp.docx")
    document.save(temp)
    with temp.open("rb") as source, OUTPUT.open("xb") as target:
        target.write(source.read())
    temp.unlink()
    print(json.dumps({"output": str(OUTPUT), "formal_l1": summary["formal_level_1_count"], "formal_l2": summary["formal_level_2_count"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
